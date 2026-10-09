# -*- coding: utf-8 -*-
"""
Excel 测试数据加载工具（数据驱动的第二种载体，与 yaml_loader 并存）
约定：第 1 行为表头（字段名），第 2 行起为数据；空行自动跳过。
"""
import os

from openpyxl import load_workbook

from utils.yaml_loader import DATA_DIR, render  # 复用 data 目录与 {ts} 模板渲染


def read_excel(file_name: str, sheet: str = "Sheet1", header_row: int = 1) -> list:
    """
    读取 data 目录下的 Excel 用例，返回 [{列名: 值}, ...]
    - data_only=True：取公式计算后的值（而不是公式本身）
    - 表头行的单元格值作为字段名，数据行与表头 zip 成字典
    """
    path = os.path.join(DATA_DIR, file_name)
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]

    headers = [cell.value for cell in ws[header_row]]
    cases = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0] is None or str(row[0]).strip() == "":  # 跳过空行
            continue
        cases.append(dict(zip(headers, row)))
    wb.close()
    return cases
